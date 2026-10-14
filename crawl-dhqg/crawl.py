#!/usr/bin/env python3
"""Cào hồ sơ SV5T cấp ĐHQG-HCM (youth.vnuhcm.edu.vn) ra một file Excel.

Đăng nhập bằng form /web/login (không qua reCAPTCHA), đọc nominee.info và các
model tiêu chí liên kết bằng JSON-RPC, rồi xuất thông tin cá nhân cùng nội dung
và file minh chứng của từng tiêu chí.
"""

from __future__ import annotations

import argparse
import getpass
import html
import http.cookiejar
import json
import re
import sys
import threading
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = "https://youth.vnuhcm.edu.vn"
HERE = Path(__file__).resolve().parent
USER_AGENT = "Mozilla/5.0"

IDENTITY_FIELDS = (
    "nominee_name",
    "nominee_id",
    "unit_id",
    "class_name",
    "nominee_type",
    "phone",
    "citizen_id",
    "gender",
    "ethnic",
    "religion",
    "academic_degree",
    "academic_title",
    "dob",
    "regular_address",
    "temp_address",
    "same_address",
    "email",
)
IDENTITY_LABELS = {
    "nominee_name": "Họ và tên",
    "nominee_id": "Mã số sinh viên/cán bộ",
    "unit_id": "Trường/Đơn vị",
    "class_name": "Khoa/Phòng/Ban",
    "nominee_type": "Nhóm đối tượng",
    "phone": "Số điện thoại",
    "citizen_id": "CMND/CCCD",
    "gender": "Giới tính",
    "ethnic": "Dân tộc",
    "religion": "Tôn giáo",
    "academic_degree": "Trình độ học vấn",
    "academic_title": "Học hàm/Học vị",
    "dob": "Ngày sinh",
    "regular_address": "Địa chỉ thường trú",
    "temp_address": "Địa chỉ tạm trú",
    "same_address": "Địa chỉ tạm trú trùng địa chỉ thường trú",
    "email": "Email",
}
SKIP_RELATIONS = {
    "ir.attachment",
    "ir.ui.view",
    "ir.ui.menu",
    "ir.model",
    "ir.model.fields",
    "ir.actions.actions",
    "ir.actions.act_window",
    "res.users",
    "res.partner",
    "res.company",
    "res.groups",
    "res.country",
    "res.currency",
    "mail.message",
    "mail.followers",
    "mail.activity",
    "mail.notification",
    "mail.thread",
    "bus.bus",
}
SKIP_FIELD_NAMES = {
    "id",
    "display_name",
    "create_uid",
    "create_date",
    "write_uid",
    "write_date",
    "__last_update",
}
TECH_PREFIXES = ("message_", "activity_", "rating_", "website_message_")
TITLE_FIELD_RE = re.compile(
    r"criter|tieu_chi|standard|indicator|question|norm|template|danh_muc|criterion",
    re.I,
)
GROUP_FIELDS = ("group_id", "parent_id", "category_id", "section_id", "standard_id")
GROUP_KEYWORDS = (
    ("tiêu chuẩn chung", 0),
    ("tieu chuan chung", 0),
    ("đạo đức", 1),
    ("dao duc", 1),
    ("học tập", 2),
    ("hoc tap", 2),
    ("thể lực", 3),
    ("the luc", 3),
    ("tình nguyện", 4),
    ("tinh nguyen", 4),
    ("hội nhập", 5),
    ("hoi nhap", 5),
)
HEADER_ALIASES = {
    "username": ("taikhoan", "tendangnhap", "email", "account", "username", "user", "mail"),
    "password": ("matkhau", "password", "pass", "pwd"),
    "student_id": ("mssv", "masinhvien", "masosinhvien", "masv", "studentid"),
    "full_name": ("hoten", "hovaten", "fullname", "name"),
}
ILLEGAL_XLSX = re.compile(r"[\000-\010\013\014\016-\037]")

_meta_lock = threading.Lock()
_field_meta: dict[str, dict] = {}
_catalog_lock = threading.Lock()
_model_links: list[ModelLink] | None = None
_catalog_cache: dict[str, list[dict]] = {}
_attachment_lock = threading.Lock()
_attachment_enabled: bool | None = None


@dataclass(frozen=True, slots=True)
class Account:
    index: int
    username: str
    password: str
    student_id: str = ""
    full_name: str = ""


@dataclass(frozen=True, slots=True)
class FileItem:
    name: str
    url: str


@dataclass(slots=True)
class Criterion:
    key: str
    group: str
    title: str
    sequence: int
    content: str = ""
    files: list[FileItem] = field(default_factory=list)
    catalog_relation: str = ""
    catalog_id: int = 0


@dataclass(frozen=True, slots=True)
class ModelLink:
    model: str
    link_field: str
    group: str
    sequence: int
    by_user: bool = False


@dataclass(slots=True)
class StudentData:
    account: Account
    scan_status: str = "Lỗi ngoài dự kiến"
    error_detail: str = ""
    profile: dict[str, str] = field(default_factory=dict)
    criteria: list[Criterion] = field(default_factory=list)


class InvalidCredentialsError(Exception):
    pass


class PortalError(Exception):
    pass


def configure_stdout() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")


def log_line(log_path: Path, status: str, email: str, message: str, *, echo: bool = False) -> None:
    text = " ".join(str(message).split())
    stamp = datetime.now().isoformat(timespec="seconds")
    line = f"{stamp}\t{status}\t{email}\t{text}"
    if echo:
        print(line, flush=True)
    with log_path.open("a", encoding="utf-8") as handle:
        handle.write(line + "\n")


def show_progress(index: int, total: int, status: str, email: str) -> None:
    width = 28
    filled = 0 if total == 0 else round(width * index / total)
    bar = "#" * filled + "-" * (width - filled)
    line = f"[{bar}] {index}/{total} {status} {email}"
    print(f"\r\x1b[2K{line[:160]}", end="", flush=True)


def finish_progress() -> None:
    print("\r\x1b[2K", end="", flush=True)


def fold(value: object) -> str:
    text = unicodedata.normalize("NFD", str(value or "").strip().lower())
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return text.replace("đ", "d")


def normalize_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", fold(value))


def column_kind(value: object) -> str | None:
    key = normalize_key(value)
    if not key:
        return None
    for kind, aliases in HEADER_ALIASES.items():
        if any(alias == key or alias in key for alias in aliases):
            return kind
    return None


def find_header(rows: list[tuple[object, ...]]) -> tuple[int, dict[str, int]]:
    candidates: list[tuple[int, dict[str, int]]] = []
    for row_index, row in enumerate(rows[:20]):
        columns: dict[str, int] = {}
        for column_index, value in enumerate(row):
            kind = column_kind(value)
            if kind and kind not in columns:
                columns[kind] = column_index
        if "username" in columns and "password" in columns:
            candidates.append((row_index, columns))
    if not candidates:
        raise ValueError("Không tìm thấy hàng tiêu đề có cả cột tài khoản và mật khẩu.")
    return max(candidates, key=lambda item: len(item[1]))


def cell_text(row: tuple[object, ...] | list[str], index: int | None) -> str:
    if index is None or index >= len(row) or row[index] is None:
        return ""
    value = row[index]
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def accounts_from_rows(rows: list[tuple[object, ...]] | list[list[str]]) -> list[Account]:
    header_index, columns = find_header([tuple(row) for row in rows])
    accounts: list[Account] = []
    for row in rows[header_index + 1 :]:
        username = cell_text(row, columns.get("username"))
        password = cell_text(row, columns.get("password"))
        if not username or not password:
            continue
        accounts.append(
            Account(
                index=len(accounts) + 1,
                username=username,
                password=password,
                student_id=cell_text(row, columns.get("student_id")),
                full_name=cell_text(row, columns.get("full_name")),
            )
        )
    return accounts


def read_accounts(file_path: Path) -> list[Account]:
    path = file_path.expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Không tìm thấy file: {path}")
    if path.suffix.lower() == ".xlsx":
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            accounts = accounts_from_rows(list(workbook.active.iter_rows(values_only=True)))
        finally:
            workbook.close()
    elif path.suffix.lower() == ".csv":
        import csv

        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            accounts = accounts_from_rows(list(csv.reader(handle)))
    else:
        raise ValueError("Chỉ hỗ trợ file .xlsx hoặc .csv.")
    if not accounts:
        raise ValueError("File không có tài khoản hợp lệ.")
    return accounts


def default_input() -> Path:
    return HERE / "data.xlsx"


class SiteSession:
    def __init__(self):
        self.opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))

    def _open(self, request):
        try:
            with self.opener.open(request, timeout=60) as response:
                return response.geturl(), response.read()
        except urllib.error.HTTPError as exc:
            raw = exc.read()
            if exc.code >= 500:
                raise PortalError(f"HTTP {exc.code}") from exc
            return exc.geturl(), raw
        except urllib.error.URLError as exc:
            raise PortalError(f"Không kết nối được web: {exc.reason}") from exc

    def get(self, url: str) -> tuple[str, str]:
        final, raw = self._open(urllib.request.Request(url, headers={"User-Agent": USER_AGENT}))
        return final, raw.decode("utf-8", "replace")

    def post_form(self, url: str, data: dict) -> tuple[str, str]:
        body = urllib.parse.urlencode(data).encode()
        request = urllib.request.Request(
            url,
            data=body,
            headers={"User-Agent": USER_AGENT, "Content-Type": "application/x-www-form-urlencoded"},
        )
        final, raw = self._open(request)
        return final, raw.decode("utf-8", "replace")

    def rpc(self, path: str, params: dict):
        payload = json.dumps({"jsonrpc": "2.0", "method": "call", "params": params, "id": 1}).encode()
        request = urllib.request.Request(
            f"{BASE}{path}",
            data=payload,
            headers={"User-Agent": USER_AGENT, "Content-Type": "application/json"},
        )
        _final, raw = self._open(request)
        try:
            body = json.loads(raw.decode())
        except json.JSONDecodeError as exc:
            raise PortalError("Web trả về dữ liệu không phải JSON") from exc
        if body.get("error"):
            data = body["error"].get("data") or {}
            message = data.get("message") or body["error"].get("message") or "Odoo lỗi"
            name = data.get("name") or ""
            if "SessionExpired" in name or "session expired" in message.lower():
                raise PortalError("Phiên đăng nhập hết hạn")
            raise PortalError(message)
        return body.get("result")


def login(username: str, password: str) -> tuple[SiteSession, dict]:
    session = SiteSession()
    _url, page = session.get(f"{BASE}/web/login")
    match = re.search(r"""name=["']csrf_token["'] value=["']([^"']+)["']""", page)
    if not match:
        raise PortalError("Không lấy được csrf đăng nhập")
    final, page = session.post_form(
        f"{BASE}/web/login",
        {
            "csrf_token": match.group(1),
            "login": username,
            "password": password,
            "redirect": "/odoo",
            "type": "password",
        },
    )
    lowered = page.lower()
    if "/web/login" in urllib.parse.urlparse(final).path or "wrong login/password" in lowered or "sai mật khẩu" in lowered or "sai tên đăng nhập" in lowered:
        raise InvalidCredentialsError("Sai tài khoản hoặc mật khẩu")
    try:
        info = session.rpc("/web/session/get_session_info", {}) or {}
    except PortalError as exc:
        if "hết hạn" in str(exc).lower():
            raise InvalidCredentialsError("Sai tài khoản hoặc mật khẩu") from exc
        raise
    uid = info.get("uid")
    signed_in = (info.get("username") or "").lower()
    if not uid or signed_in != username.lower():
        raise InvalidCredentialsError("Không đăng nhập được")
    return session, info


def call_kw(session: SiteSession, model: str, method: str, args, kwargs=None):
    return session.rpc(
        f"/web/dataset/call_kw/{model}/{method}",
        {"model": model, "method": method, "args": args, "kwargs": kwargs or {}},
    )


def chunks(items: list, size: int):
    for start in range(0, len(items), size):
        yield items[start : start + size]


def is_access_error(message: str) -> bool:
    lowered = message.lower()
    return "access" in lowered or "not allowed" in lowered or "quyền" in lowered or "truy cập" in lowered


def fields_meta(session: SiteSession, model: str) -> dict:
    with _meta_lock:
        cached = _field_meta.get(model)
    if cached is not None:
        return cached
    raw = call_kw(
        session,
        model,
        "fields_get",
        [],
        {"attributes": ["string", "type", "relation", "selection"], "context": {"lang": "vi_VN"}},
    )
    meta = raw or {}
    with _meta_lock:
        _field_meta[model] = meta
    return meta


def readable_fields(meta: dict) -> list[str]:
    names = ["id", "display_name"]
    for name, spec in meta.items():
        if name in names or name in SKIP_FIELD_NAMES or name.startswith(TECH_PREFIXES):
            continue
        if spec.get("type") in {"one2many", "many2many", "binary", "image"} or name.startswith("image"):
            names.append(name)
            continue
        names.append(name)
    return names


def _without_blocked_fields(fields: list[str], message: str) -> list[str]:
    blocked = set(re.findall(r"['\"]([A-Za-z_][A-Za-z0-9_]*)['\"]", message))
    reduced = [name for name in fields if name not in blocked and not name.startswith("image")]
    return reduced if reduced and reduced != fields else []


def search_read(session: SiteSession, model: str, domain, fields: list[str], limit: int = 500) -> list[dict]:
    kwargs = {"fields": fields, "limit": limit, "context": {"bin_size": True, "lang": "vi_VN"}}
    try:
        rows = call_kw(session, model, "search_read", [domain], kwargs)
    except PortalError as exc:
        reduced = _without_blocked_fields(fields, str(exc))
        if not reduced:
            raise
        kwargs["fields"] = reduced
        rows = call_kw(session, model, "search_read", [domain], kwargs)
    return rows or []


def read_ids(session: SiteSession, model: str, ids: list[int], fields: list[str]) -> list[dict]:
    rows: list[dict] = []
    active_fields = fields
    for chunk in chunks(ids, 80):
        try:
            found = call_kw(
                session,
                model,
                "read",
                [chunk, active_fields],
                {"context": {"bin_size": True, "lang": "vi_VN"}},
            )
        except PortalError as exc:
            reduced = _without_blocked_fields(active_fields, str(exc))
            if not reduced:
                raise
            active_fields = reduced
            found = call_kw(
                session,
                model,
                "read",
                [chunk, active_fields],
                {"context": {"bin_size": True, "lang": "vi_VN"}},
            )
        rows.extend(found or [])
    return rows


def attachment_items(session: SiteSession, ids: list[int]) -> list[FileItem]:
    clean = []
    for item in ids:
        if isinstance(item, (list, tuple)) and item:
            clean.append(int(item[0]))
        elif isinstance(item, int):
            clean.append(item)
    if not clean:
        return []
    names: dict[int, str] = {}
    try:
        rows = read_ids(session, "ir.attachment", clean, ["name"])
        names = {int(row["id"]): str(row.get("name") or f"file-{row['id']}") for row in rows}
    except PortalError:
        names = {item: f"file-{item}" for item in clean}
    return [FileItem(names.get(item, f"file-{item}"), f"{BASE}/web/content/{item}?download=true") for item in clean]


def attachments_for_records(session: SiteSession, model: str, ids: list[int]) -> dict[int, list[FileItem]]:
    global _attachment_enabled
    with _attachment_lock:
        enabled = _attachment_enabled
    if enabled is False or not ids:
        return {}
    try:
        rows = search_read(
            session,
            "ir.attachment",
            [["res_model", "=", model], ["res_id", "in", ids], ["type", "=", "binary"]],
            ["name", "res_id"],
            limit=1000,
        )
    except PortalError as exc:
        if is_access_error(str(exc)):
            with _attachment_lock:
                _attachment_enabled = False
            return {}
        return {}
    with _attachment_lock:
        _attachment_enabled = True
    grouped: dict[int, list[FileItem]] = defaultdict(list)
    for row in rows:
        res_id = int(row.get("res_id") or 0)
        if not res_id:
            continue
        grouped[res_id].append(
            FileItem(str(row.get("name") or f"file-{row['id']}"), f"{BASE}/web/content/{row['id']}?download=true")
        )
    return grouped


def relation_ids(value) -> list[int]:
    if not value or value is False:
        return []
    if isinstance(value, int):
        return [value]
    ids = []
    for item in value:
        if isinstance(item, int):
            ids.append(item)
        elif isinstance(item, (list, tuple)) and item and isinstance(item[0], int):
            ids.append(item[0])
    return ids


def is_line_field(name: str, spec: dict) -> bool:
    if name.startswith(TECH_PREFIXES) or name in SKIP_FIELD_NAMES:
        return False
    if spec.get("type") not in {"one2many", "many2many"}:
        return False
    relation = spec.get("relation") or ""
    if not relation or relation in SKIP_RELATIONS or relation.startswith(("mail.", "bus.", "ir.")):
        return False
    return True


def dedupe_files(items: list[FileItem]) -> list[FileItem]:
    seen = set()
    result = []
    for item in items:
        marker = (item.name, item.url)
        if marker in seen:
            continue
        seen.add(marker)
        result.append(item)
    return result


def files_on_row(session: SiteSession, model: str, row: dict, meta: dict) -> list[FileItem]:
    files: list[FileItem] = []
    for name, spec in meta.items():
        if spec.get("relation") != "ir.attachment":
            if spec.get("type") == "binary" and not name.startswith("image") and row.get(name):
                files.append(
                    FileItem(
                        spec.get("string") or name,
                        f"{BASE}/web/content?model={urllib.parse.quote(model)}&id={row['id']}&field={urllib.parse.quote(name)}&download=true",
                    )
                )
            continue
        files.extend(attachment_items(session, relation_ids(row.get(name))))
    return dedupe_files(files)


def hydrate(session: SiteSession, model: str, ids: list[int], depth: int = 0) -> list[dict]:
    if not ids or depth > 3:
        return []
    meta = fields_meta(session, model)
    rows = read_ids(session, model, ids, readable_fields(meta))
    attached = attachments_for_records(session, model, [int(row["id"]) for row in rows])
    for row in rows:
        row["_model"] = model
        row["_meta"] = meta
        row["_files"] = dedupe_files(files_on_row(session, model, row, meta) + attached.get(int(row["id"]), []))
        row["_children"] = {}
        if depth >= 2:
            continue
        for name, spec in meta.items():
            if not is_line_field(name, spec):
                continue
            child_ids = relation_ids(row.get(name))
            if child_ids:
                row["_children"][name] = hydrate(session, spec["relation"], child_ids, depth + 1)
    return rows


def plain_text(value: str) -> str:
    text = re.sub(r"(?i)<br\s*/?>", "\n", value)
    text = re.sub(r"(?i)</p>", "\n", text)
    text = re.sub(r"<[^>]+>", "", text)
    text = html.unescape(text)
    text = text.replace("\xa0", " ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def selection_label(spec: dict, value) -> str:
    for item in spec.get("selection") or []:
        if isinstance(item, (list, tuple)) and len(item) >= 2 and item[0] == value:
            return str(item[1])
    return "" if value in (None, False) else str(value)


def format_scalar(value, spec: dict) -> str:
    kind = spec.get("type")
    if kind == "boolean":
        if value is True:
            return "Có"
        if value is False:
            return "Không"
        return ""
    if value in (None, False, ""):
        return ""
    if kind == "selection":
        return selection_label(spec, value).strip()
    if kind == "many2one":
        if isinstance(value, (list, tuple)) and len(value) > 1:
            return str(value[1]).strip()
        return ""
    if kind in {"html", "text"}:
        return plain_text(str(value))
    if kind in {"integer", "float", "monetary"}:
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()
    if isinstance(value, (list, tuple, dict)):
        return ""
    return plain_text(str(value)) if "<" in str(value) else str(value).strip()


def field_label(name: str, spec: dict) -> str:
    if name in IDENTITY_LABELS:
        return IDENTITY_LABELS[name]
    label = str(spec.get("string") or name).strip()
    return re.sub(r"\s+", " ", label)


def image_url(record_id: int, meta: dict, record: dict) -> str:
    for name in ("image_original", "image_1920", "image_512", "avatar"):
        spec = meta.get(name) or {}
        if spec.get("type") in {"binary", "image"} and record.get(name):
            return f"{BASE}/web/image/nominee.info/{record_id}/{name}"
    return ""


def shape_profile(record: dict, meta: dict) -> tuple[dict[str, str], list[Criterion]]:
    profile: dict[str, str] = {}
    extras: list[Criterion] = []
    for name in IDENTITY_FIELDS:
        spec = meta.get(name)
        if not spec or spec.get("type") in {"one2many", "many2many", "binary", "image"}:
            continue
        profile[field_label(name, spec)] = format_scalar(record.get(name), spec)
    avatar = image_url(int(record["id"]), meta, record)
    profile["Ảnh đại diện"] = "Đã có" if avatar else "Chưa có"
    profile["Ảnh đại diện (Link)"] = avatar
    profile_files = dedupe_files(record.get("_files") or [])
    if profile_files:
        profile["Tệp đính kèm hồ sơ"] = "\n".join(item.name for item in profile_files)
        profile["Tệp đính kèm hồ sơ (Link)"] = "\n".join(item.url for item in profile_files if item.url)
    sequence = 0
    for name, spec in meta.items():
        if name in IDENTITY_FIELDS or name in SKIP_FIELD_NAMES or name.startswith(TECH_PREFIXES):
            continue
        if name.startswith("image") or spec.get("type") in {"one2many", "many2many", "binary", "image"}:
            continue
        value = format_scalar(record.get(name), spec)
        if not value or (spec.get("type") == "boolean" and value == "Không"):
            continue
        label = field_label(name, spec)
        if len(label) < 90:
            profile.setdefault(label, value)
            continue
        sequence += 1
        extras.append(
            Criterion(
                key=f"field:{name}",
                group="NỘI DUNG TRÊN HỒ SƠ",
                title=label,
                sequence=sequence,
                content=value,
            )
        )
    return profile, extras


def title_source(record: dict, meta: dict) -> tuple[str, str, int]:
    for name, spec in meta.items():
        if spec.get("type") != "many2one":
            continue
        relation = spec.get("relation") or ""
        if not TITLE_FIELD_RE.search(f"{name} {relation}"):
            continue
        value = record.get(name)
        if isinstance(value, (list, tuple)) and len(value) > 1 and value[1]:
            return str(value[1]).strip(), relation, int(value[0])
    for name in ("name", "title", "code"):
        spec = meta.get(name) or {}
        if spec.get("type") not in {"char", "text", "selection"}:
            continue
        text = format_scalar(record.get(name), spec)
        if text and len(text) <= 250 and "\n" not in text:
            return text, "", 0
    display = str(record.get("display_name") or "").strip()
    return display or "Mục", "", 0


def content_lines(record: dict, meta: dict, title: str) -> list[str]:
    lines = []
    for name, spec in meta.items():
        if name in SKIP_FIELD_NAMES or name.startswith(TECH_PREFIXES) or name in {"sequence", "color"}:
            continue
        if spec.get("type") in {"one2many", "many2many", "binary", "image"}:
            continue
        if name.startswith("image"):
            continue
        text = format_scalar(record.get(name), spec)
        if not text:
            continue
        if spec.get("type") == "boolean" and text == "Không":
            continue
        label = field_label(name, spec)
        if text == title:
            continue
        if spec.get("type") in {"char", "text", "html"} and label.lower() in {"tên", "tiêu đề", "name", "title"}:
            continue
        if spec.get("type") in {"text", "html", "char"} and len([item for item in lines if item]) == 0 and text != label:
            lines.append(text)
            continue
        lines.append(f"{label}: {text}" if text != label else text)
    return [line for line in lines if line]


def looks_like_section(record: dict) -> bool:
    children = {key: value for key, value in (record.get("_children") or {}).items() if value}
    if len(children) != 1:
        return False
    meta = record.get("_meta") or {}
    for name, spec in meta.items():
        if spec.get("type") not in {"text", "html"}:
            continue
        text = format_scalar(record.get(name), spec)
        if len(text) > 40:
            return False
    return True


def line_sequence(record: dict, fallback: int) -> int:
    for name in ("sequence", "stt", "order", "index"):
        value = record.get(name)
        if isinstance(value, int):
            return value
        if isinstance(value, float) and value.is_integer():
            return int(value)
    return fallback


def rows_to_criteria(rows: list[dict], default_group: str, sequence_base: int) -> list[Criterion]:
    criteria: list[Criterion] = []
    for index, record in enumerate(rows):
        meta = record.get("_meta") or {}
        if looks_like_section(record):
            child_name = next(name for name, value in record["_children"].items() if value)
            title, _relation, _catalog_id = title_source(record, meta)
            group = title or default_group
            nested = rows_to_criteria(record["_children"][child_name], group, sequence_base + index * 1000)
            extra_files = record.get("_files") or []
            if extra_files:
                for item in nested:
                    item.files = dedupe_files(list(item.files) + list(extra_files))
            criteria.extend(nested)
            continue
        title, relation, catalog_id = title_source(record, meta)
        parts = content_lines(record, meta, title)
        files = list(record.get("_files") or [])
        for child_rows in (record.get("_children") or {}).values():
            for child in child_rows:
                child_meta = child.get("_meta") or {}
                child_title, _rel, _cid = title_source(child, child_meta)
                child_parts = content_lines(child, child_meta, child_title)
                if child_parts:
                    parts.append(child_title + ": " + "\n".join(child_parts) if child_title and child_title != "Mục" else "\n".join(child_parts))
                files.extend(child.get("_files") or [])
        key = f"cat:{relation}:{catalog_id}" if catalog_id else f"line:{fold(default_group)}:{fold(title)}"
        criteria.append(
            Criterion(
                key=key,
                group=default_group or "TIÊU CHÍ",
                title=title or "Mục",
                sequence=sequence_base + line_sequence(record, index),
                content="\n".join(parts).strip(),
                files=dedupe_files(files),
                catalog_relation=relation,
                catalog_id=catalog_id,
            )
        )
    return criteria


def catalog_group(item: dict) -> str:
    for name in GROUP_FIELDS:
        value = item.get(name)
        if isinstance(value, (list, tuple)) and len(value) > 1 and value[1]:
            return str(value[1]).strip()
    return ""


def catalog_title(item: dict) -> str:
    for name in ("name", "title", "display_name", "code"):
        value = item.get(name)
        if value and value is not False:
            text = str(value).strip()
            if text:
                return text
    return "Mục"


def catalog_sequence(item: dict, fallback: int) -> int:
    value = item.get("sequence")
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return fallback


def merge_catalog(criteria: list[Criterion], catalogs: dict[str, list[dict]]) -> list[Criterion]:
    merged: dict[str, Criterion] = {}
    order: list[str] = []
    for item in criteria:
        if item.key not in merged:
            merged[item.key] = item
            order.append(item.key)
            continue
        current = merged[item.key]
        if item.content and item.content not in current.content:
            current.content = "\n".join(part for part in (current.content, item.content) if part)
        current.files = dedupe_files(current.files + item.files)
    for relation, rows in catalogs.items():
        for index, row in enumerate(rows):
            catalog_id = int(row.get("id") or 0)
            if not catalog_id:
                continue
            key = f"cat:{relation}:{catalog_id}"
            group = catalog_group(row)
            title = catalog_title(row)
            if key in merged:
                current = merged[key]
                if group:
                    current.group = group
                if title and title != "Mục":
                    current.title = title
                current.sequence = catalog_sequence(row, current.sequence)
                continue
            merged[key] = Criterion(
                key=key,
                group=group or "TIÊU CHÍ",
                title=title,
                sequence=catalog_sequence(row, 10000 + index),
                catalog_relation=relation,
                catalog_id=catalog_id,
            )
            order.append(key)
    return [merged[key] for key in order]


def title_relations_of(session: SiteSession, model: str, depth: int = 0, seen: set[str] | None = None) -> list[str]:
    seen = set() if seen is None else seen
    if model in seen:
        return []
    seen.add(model)
    try:
        meta = fields_meta(session, model)
    except PortalError:
        return []
    found: list[str] = []
    for name, spec in meta.items():
        relation = spec.get("relation") or ""
        if spec.get("type") == "many2one" and relation and TITLE_FIELD_RE.search(f"{name} {relation}"):
            found.append(relation)
        elif depth < 2 and is_line_field(name, spec):
            found.extend(title_relations_of(session, relation, depth + 1, seen))
    return found


def load_catalog(session: SiteSession, relation: str) -> list[dict]:
    with _catalog_lock:
        if relation in _catalog_cache:
            return _catalog_cache[relation]
    try:
        meta = fields_meta(session, relation)
    except PortalError:
        with _catalog_lock:
            _catalog_cache[relation] = []
        return []
    wanted = [name for name in ("display_name", "name", "title", "code", "sequence", *GROUP_FIELDS) if name in meta]
    if "id" not in wanted:
        wanted.insert(0, "id")
    try:
        rows = search_read(session, relation, [], wanted, limit=400)
    except PortalError:
        rows = []
    if not rows or len(rows) > 350:
        rows = []
    with _catalog_lock:
        _catalog_cache[relation] = rows
    return rows


def load_menus(session: SiteSession) -> dict:
    try:
        _url, raw = session.get(f"{BASE}/web/webclient/load_menus/1")
        data = json.loads(raw)
    except (PortalError, json.JSONDecodeError, urllib.error.URLError):
        return {}
    return data if isinstance(data, dict) else {}


def menu_actions(menus: dict) -> list[tuple[int, str, int]]:
    found: list[tuple[int, str, int]] = []

    def visit(node: dict, parent_name: str, order: int) -> None:
        name = str(node.get("name") or parent_name or "").strip()
        action = node.get("action")
        action_id = 0
        if isinstance(action, str) and "," in action:
            tail = action.split(",")[-1]
            if tail.isdigit():
                action_id = int(tail)
        elif isinstance(action, int):
            action_id = action
        if action_id:
            found.append((action_id, name or parent_name, order))
        children = node.get("children") or []
        for child_index, child in enumerate(children):
            if isinstance(child, int):
                child = menus.get(str(child)) or menus.get(child)
            if isinstance(child, dict):
                visit(child, name, order * 100 + child_index)

    root = menus.get("root")
    if isinstance(root, dict):
        visit(root, "", 1)
    return found


def discover_links(session: SiteSession) -> list[ModelLink]:
    global _model_links
    with _catalog_lock:
        if _model_links is not None:
            return _model_links
    links: list[ModelLink] = []
    seen = set()
    for action_id, menu_name, order in menu_actions(load_menus(session)):
        try:
            action = session.rpc("/web/action/load", {"action_id": action_id}) or {}
        except PortalError:
            continue
        model = action.get("res_model") or ""
        if not model or model == "nominee.info" or model in SKIP_RELATIONS or model.startswith(("mail.", "ir.", "bus.", "res.")):
            continue
        try:
            meta = fields_meta(session, model)
        except PortalError:
            continue
        link_field = ""
        by_user = False
        for preferred in ("nominee_id", "nominee_info_id", "info_id", "profile_id"):
            spec = meta.get(preferred) or {}
            if spec.get("type") == "many2one" and spec.get("relation") == "nominee.info":
                link_field = preferred
                break
        if not link_field:
            for name, spec in meta.items():
                if spec.get("type") == "many2one" and spec.get("relation") == "nominee.info":
                    link_field = name
                    break
        if not link_field and (meta.get("user_id") or {}).get("relation") == "res.users":
            link_field = "user_id"
            by_user = True
        if not link_field or (model, link_field) in seen:
            continue
        seen.add((model, link_field))
        links.append(ModelLink(model, link_field, menu_name or model, order, by_user))
    with _catalog_lock:
        _model_links = links
    return links


def find_nominee(session: SiteSession, uid: int) -> tuple[dict | None, str]:
    meta = fields_meta(session, "nominee.info")
    fields = readable_fields(meta)
    user_fields = [
        name
        for name, spec in meta.items()
        if spec.get("type") == "many2one" and spec.get("relation") == "res.users"
    ]
    ordered = [name for name in ("user_id", "create_uid") if name in user_fields]
    ordered += [name for name in user_fields if name not in ordered]
    note = ""
    for name in ordered:
        rows = search_read(session, "nominee.info", [[name, "=", uid]], fields, limit=10)
        if not rows:
            continue
        chosen = max(rows, key=lambda row: int(row["id"]))
        if len(rows) > 1:
            note = f"Có {len(rows)} hồ sơ, lấy id={chosen['id']}"
        return chosen, note
    return None, note


def fetch_account(session: SiteSession, info: dict) -> tuple[dict[str, str], list[Criterion], str, bool]:
    nominee, note = find_nominee(session, int(info["uid"]))
    if not nominee:
        return {}, [], note, False
    meta = fields_meta(session, "nominee.info")
    hydrated = hydrate(session, "nominee.info", [int(nominee["id"])])
    record = hydrated[0] if hydrated else nominee
    record.setdefault("_meta", meta)
    record.setdefault("_files", [])
    record.setdefault("_children", {})
    profile, criteria = shape_profile(record, meta)
    warnings: list[str] = []
    seen_records: set[tuple[str, int]] = set()

    def claim(rows: list[dict]) -> list[dict]:
        fresh = []
        for row in rows:
            marker = (str(row.get("_model") or ""), int(row.get("id") or 0))
            if marker[1] and marker in seen_records:
                continue
            if marker[1]:
                seen_records.add(marker)
            fresh.append(row)
            for nested in (row.get("_children") or {}).values():
                claim(nested)
        return fresh

    sequence = 100
    for name, spec in meta.items():
        if not is_line_field(name, spec):
            continue
        children = claim((record.get("_children") or {}).get(name) or [])
        group = field_label(name, spec) or "TIÊU CHÍ"
        try:
            criteria.extend(rows_to_criteria(children, group, sequence))
        except PortalError as exc:
            warnings.append(f"{group}: {exc}")
        sequence += 1000
    for link in discover_links(session):
        try:
            if link.by_user:
                domain = [[link.link_field, "=", int(info["uid"])]]
            else:
                domain = [[link.link_field, "=", int(record["id"])]]
            link_meta = fields_meta(session, link.model)
            found = search_read(session, link.model, domain, ["id"], limit=500)
            ids = [int(row["id"]) for row in found]
            if not ids:
                continue
            rows = claim(hydrate(session, link.model, ids))
            for row in rows:
                row.setdefault("_meta", link_meta)
            criteria.extend(rows_to_criteria(rows, link.group or link.model, link.sequence))
        except PortalError as exc:
            warnings.append(f"{link.group or link.model}: {exc}")
    relations = {item.catalog_relation for item in criteria if item.catalog_relation}
    relations.update(title_relations_of(session, "nominee.info"))
    for link in discover_links(session):
        relations.update(title_relations_of(session, link.model))
    catalogs = {relation: load_catalog(session, relation) for relation in sorted(relations)}
    criteria = merge_catalog(criteria, catalogs)
    detail = "; ".join(part for part in ([note] if note else []) + warnings if part)
    return profile, criteria, detail, bool(warnings)


def scan_account(account: Account, retries: int) -> StudentData:
    last_error = ""
    for attempt in range(retries + 1):
        try:
            session, info = login(account.username, account.password)
            profile, criteria, note, partial = fetch_account(session, info)
            student = StudentData(account)
            student.profile = profile
            student.criteria = criteria
            student.error_detail = note
            if not profile and not criteria:
                student.scan_status = "Chưa có hồ sơ"
                student.error_detail = note or "Đã đăng nhập nhưng chưa có hồ sơ cá nhân"
            elif partial:
                student.scan_status = "Lỗi một phần"
            else:
                student.scan_status = "Thành công"
            return student
        except InvalidCredentialsError as exc:
            student = StudentData(account)
            student.scan_status = "Sai tài khoản/mật khẩu"
            student.error_detail = str(exc)
            return student
        except Exception as exc:
            last_error = str(exc)
            if attempt >= retries:
                break
            time.sleep(1.5 * (attempt + 1))
    student = StudentData(account)
    student.scan_status = "Lỗi ngoài dự kiến"
    student.error_detail = f"{last_error} (đã thử {retries + 1} lần)"
    return student


def group_rank(name: str) -> int:
    folded = fold(name)
    for keyword, rank in GROUP_KEYWORDS:
        if keyword in folded:
            return rank
    return 50


def ordered_criteria(results: list[StudentData]) -> list[Criterion]:
    catalog: dict[str, Criterion] = {}
    for result in results:
        for item in result.criteria:
            current = catalog.get(item.key)
            if current is None or (item.sequence, item.title) < (current.sequence, current.title):
                catalog[item.key] = item
    return sorted(catalog.values(), key=lambda item: (group_rank(item.group), fold(item.group), item.sequence, fold(item.title)))


def profile_columns(results: list[StudentData]) -> list[str]:
    preferred = [IDENTITY_LABELS[name] for name in IDENTITY_FIELDS]
    preferred += ["Ảnh đại diện", "Ảnh đại diện (Link)", "Tệp đính kèm hồ sơ", "Tệp đính kèm hồ sơ (Link)"]
    columns = []
    seen = set()
    present = {key for result in results for key in result.profile}
    for label in preferred:
        if label in present and label not in seen:
            columns.append(label)
            seen.add(label)
    extras = sorted(present - seen, key=fold)
    return columns + extras


def safe_cell(value: object) -> str:
    text = ILLEGAL_XLSX.sub("", "" if value is None else str(value))
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def apply_header(cell, fill, font, border) -> None:
    cell.fill = fill
    cell.font = font
    cell.border = border
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_excel(results: list[StudentData], output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Hồ sơ SV5T ĐHQG"
    sheet.sheet_view.showGridLines = True
    sheet.views.sheetView[0].zoomScale = 90

    base_headers = ["STT", "MSSV", "Họ và tên (File)", "Tài khoản (File)", "Trạng thái quét", "Chi tiết lỗi"]
    profiles = profile_columns(results)
    criteria = ordered_criteria(results)
    grouped: list[tuple[str, list[Criterion]]] = []
    for item in criteria:
        if grouped and grouped[-1][0] == item.group:
            grouped[-1][1].append(item)
        else:
            grouped.append((item.group, [item]))

    fill_account = PatternFill("solid", fgColor="1F4E78")
    fill_profile = PatternFill("solid", fgColor="203764")
    fill_group = PatternFill("solid", fgColor="2F5597")
    fill_sub_account = PatternFill("solid", fgColor="D9E1F2")
    fill_sub_profile = PatternFill("solid", fgColor="DDEBF7")
    fill_sub_group = PatternFill("solid", fgColor="EAEEF3")
    fill_detail = PatternFill("solid", fgColor="F2F4F7")
    font_top = Font(name="Arial", size=10.5, bold=True, color="FFFFFF")
    font_mid = Font(name="Arial", size=9, bold=True, color="1F1F1F")
    font_low = Font(name="Arial", size=8.5, bold=True, color="333333")
    border = Border(
        left=Side(style="thin", color="B0C4DE"),
        right=Side(style="thin", color="B0C4DE"),
        top=Side(style="thin", color="B0C4DE"),
        bottom=Side(style="thin", color="B0C4DE"),
    )
    widths: dict[int, int] = {}
    column = 1

    def paint_band(start: int, end: int, title: str, top_fill, sub_fill, labels: list[str], merge_down: bool) -> None:
        if end < start:
            return
        sheet.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
        for index in range(start, end + 1):
            top = sheet.cell(1, index)
            top.fill = top_fill
            top.border = border
            top.font = font_top
            top.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if merge_down:
                sheet.merge_cells(start_row=2, start_column=index, end_row=3, end_column=index)
            label = labels[index - start] if index - start < len(labels) else ""
            mid = sheet.cell(2, index, label if merge_down else None)
            mid.fill = sub_fill
            mid.font = font_mid
            mid.border = border
            mid.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            low = sheet.cell(3, index)
            low.fill = sub_fill
            low.border = border
        sheet.cell(1, start, title)

    paint_band(column, column + len(base_headers) - 1, "THÔNG TIN TÀI KHOẢN & TRẠNG THÁI", fill_account, fill_sub_account, base_headers, True)
    for offset, _label in enumerate(base_headers):
        widths[column + offset] = (6, 16, 28, 32, 22, 36)[offset]
    column += len(base_headers)

    if profiles:
        paint_band(column, column + len(profiles) - 1, "THÔNG TIN CÁ NHÂN", fill_profile, fill_sub_profile, profiles, True)
        for offset in range(len(profiles)):
            widths[column + offset] = 28
        column += len(profiles)

    detail_labels = ["Nội dung / Lựa chọn", "Tên tệp đính kèm", "Link tải tệp"]
    for group_name, items in grouped:
        start = column
        end = column + len(items) * 3 - 1
        sheet.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
        for index in range(start, end + 1):
            cell = sheet.cell(1, index)
            cell.fill = fill_group
            cell.border = border
            cell.font = font_top
        sheet.cell(1, start, group_name.upper())
        sheet.cell(1, start).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for item in items:
            sheet.merge_cells(start_row=2, start_column=column, end_row=2, end_column=column + 2)
            for offset in range(3):
                cell = sheet.cell(2, column + offset)
                cell.fill = fill_sub_group
                cell.border = border
                cell.font = font_mid
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            sheet.cell(2, column, item.title)
            for offset, label in enumerate(detail_labels):
                cell = sheet.cell(3, column + offset, label)
                apply_header(cell, fill_detail, font_low, border)
                widths[column + offset] = (42, 32, 36)[offset]
            column += 3

    total_columns = max(column - 1, 1)
    sheet.row_dimensions[1].height = 34
    sheet.row_dimensions[2].height = 48
    sheet.row_dimensions[3].height = 26
    for index, width in widths.items():
        sheet.column_dimensions[get_column_letter(index)].width = width

    body_font = Font(name="Arial", size=9)
    body_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )
    zebra = PatternFill("solid", fgColor="F9FAFB")
    by_key = [item.key for item in criteria]
    for row_index, result in enumerate(results):
        target = row_index + 4
        values = [
            str(result.account.index),
            result.account.student_id,
            result.profile.get("Họ và tên") or result.account.full_name,
            result.account.username,
            result.scan_status,
            result.error_detail,
        ]
        for label in profiles:
            values.append(result.profile.get(label, ""))
        found = {item.key: item for item in result.criteria}
        for key in by_key:
            item = found.get(key)
            if item:
                values.append(item.content)
                values.append("\n".join(file.name for file in item.files))
                values.append("\n".join(file.url for file in item.files if file.url))
            else:
                values.extend(["", "", ""])
        for col_index, value in enumerate(values, start=1):
            cell = sheet.cell(target, col_index, safe_cell(value))
            cell.font = body_font
            cell.border = body_border
            if row_index % 2 == 1:
                cell.fill = zebra
            cell.alignment = Alignment(
                horizontal="center" if col_index in {1, 2, 5} else "left",
                vertical="center",
                wrap_text=True,
            )
        sheet.row_dimensions[target].height = 36

    sheet.freeze_panes = "E4"
    if results:
        sheet.auto_filter.ref = f"A3:{get_column_letter(total_columns)}{len(results) + 3}"
    sheet.auto_filter.ref = sheet.auto_filter.ref or f"A3:{get_column_letter(total_columns)}3"
    temporary = output_path.with_name(f".{output_path.stem}.tmp.xlsx")
    try:
        workbook.save(temporary)
        temporary.replace(output_path)
    finally:
        if temporary.exists() and temporary != output_path:
            temporary.unlink(missing_ok=True)
    return output_path


def write_schema(path: Path, links: list[ModelLink]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    models = {}
    for model, meta in sorted(_field_meta.items()):
        models[model] = {
            name: {
                "type": spec.get("type"),
                "string": spec.get("string"),
                "relation": spec.get("relation") or "",
            }
            for name, spec in meta.items()
            if not name.startswith(TECH_PREFIXES)
        }
    payload = {
        "models": models,
        "links": [
            {"model": link.model, "field": link.link_field, "menu": link.group, "by_user": link.by_user}
            for link in links
        ],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def scan_all(accounts: list[Account], workers: int, retries: int, log_path: Path) -> list[StudentData]:
    results: list[StudentData] = []
    lock = threading.Lock()
    done = 0
    total = len(accounts)

    def worker(account: Account) -> StudentData:
        nonlocal done
        student = scan_account(account, retries)
        with lock:
            done += 1
            log_line(log_path, student.scan_status, account.username, student.error_detail or f"{len(student.criteria)} tiêu chí")
            show_progress(done, total, student.scan_status, account.username)
        return student

    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        futures = [pool.submit(worker, account) for account in accounts]
        for future in as_completed(futures):
            results.append(future.result())
    finish_progress()
    return sorted(results, key=lambda item: item.account.index)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Cào hồ sơ và tiêu chí SV5T ĐHQG-HCM ra Excel.")
    parser.add_argument("input_file", nargs="?", help="File tài khoản. Mặc định: crawl-dhqg/data.xlsx.")
    parser.add_argument("--account", help="Cào một tài khoản. Mật khẩu được hỏi lúc chạy.")
    parser.add_argument("--output", type=Path, help="File Excel đầu ra")
    parser.add_argument("--workers", type=int, default=6, help="Số luồng đăng nhập song song (mặc định 6)")
    parser.add_argument("--limit", type=int, default=0, help="Chỉ cào N tài khoản đầu")
    parser.add_argument("--retries", type=int, default=2, help="Số lần thử lại khi lỗi mạng (mặc định 2)")
    parser.add_argument("--self-test", action="store_true", help=argparse.SUPPRESS)
    return parser.parse_args()


def self_test() -> int:
    meta = {
        "name": {"type": "char", "string": "Tên"},
        "note": {"type": "text", "string": "Minh chứng"},
        "passed": {"type": "boolean", "string": "Đạt"},
        "criterion_id": {"type": "many2one", "string": "Tiêu chí", "relation": "youth.criterion"},
    }
    line = {
        "id": 7,
        "criterion_id": [3, "Đạt học lực giỏi"],
        "name": False,
        "note": "Điểm trung bình 8.5",
        "passed": True,
        "_meta": meta,
        "_files": [FileItem("bang-diem.pdf", f"{BASE}/web/content/9?download=true")],
        "_children": {},
    }
    section_meta = {"name": {"type": "char", "string": "Tên"}, "line_ids": {"type": "one2many", "string": "Chi tiết", "relation": "youth.line"}}
    section = {
        "id": 1,
        "name": "Học tập tốt",
        "_meta": section_meta,
        "_files": [],
        "_children": {"line_ids": [line]},
    }
    nominee_meta = {
        "nominee_name": {"type": "char", "string": "Name"},
        "nominee_id": {"type": "char", "string": "Student code"},
        "gender": {"type": "selection", "string": "Gender", "selection": [["male", "Nam"], ["female", "Nữ"]]},
        "line_ids": {"type": "one2many", "string": "Tiêu chí", "relation": "youth.section"},
    }
    nominee_meta["image_1920"] = {"type": "binary", "string": "Ảnh"}
    nominee = {
        "id": 4,
        "nominee_name": "Nguyễn Văn A",
        "nominee_id": "2210001",
        "gender": "male",
        "image_1920": "12.0 Kb",
        "_files": [],
        "_children": {},
        "_meta": nominee_meta,
    }
    profile, extras = shape_profile(nominee, nominee_meta)
    criteria = extras + rows_to_criteria([section], "Tiêu chí", 100)
    criteria = merge_catalog(
        criteria,
        {
            "youth.criterion": [
                {"id": 3, "name": "Đạt học lực giỏi", "sequence": 2, "group_id": [1, "Học tập tốt"]},
                {"id": 4, "name": "Không nợ môn", "sequence": 3, "group_id": [1, "Học tập tốt"]},
            ]
        },
    )
    assert profile["Họ và tên"] == "Nguyễn Văn A", profile
    assert profile["Giới tính"] == "Nam", profile
    assert profile["Ảnh đại diện"] == "Đã có"
    titles = [(item.group, item.title, item.content) for item in criteria]
    assert ("Học tập tốt", "Đạt học lực giỏi", "Điểm trung bình 8.5\nĐạt: Có") == titles[0], titles
    assert any(item.title == "Không nợ môn" and item.content == "" for item in criteria), titles
    account = Account(1, "a@example.com", "secret", "2210001", "Nguyễn Văn A")
    student = StudentData(account, "Thành công", "", profile, criteria)
    output = HERE / ".self-test.xlsx"
    build_excel([student], output)
    workbook = openpyxl.load_workbook(output, read_only=True, data_only=True)
    saved = list(workbook.active.iter_rows(values_only=True))
    workbook.close()
    output.unlink(missing_ok=True)
    assert saved[0][0] == "THÔNG TIN TÀI KHOẢN & TRẠNG THÁI"
    assert "HỌC TẬP TỐT" in saved[0]
    assert "Đạt học lực giỏi" in saved[1]
    assert "Điểm trung bình 8.5" in " ".join(str(cell or "") for cell in saved[3])
    print("self-test ok")
    return 0


def main() -> int:
    configure_stdout()
    args = parse_args()
    if args.self_test:
        return self_test()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output = args.output or (HERE / "kết quả" / f"Du_lieu_SV5T_DHQG_{timestamp}.xlsx")
    log_path = HERE / "nhật ký" / f"Nhat_ky_SV5T_DHQG_{timestamp}.log"
    log_path.parent.mkdir(parents=True, exist_ok=True)

    if args.account:
        password = getpass.getpass("Mật khẩu: ")
        accounts = [Account(1, args.account.strip(), password)]
        source = args.account.strip()
    else:
        chosen = Path(args.input_file) if args.input_file else default_input()
        if not chosen.is_file():
            print(f"Không thấy file tài khoản: {chosen}", file=sys.stderr)
            return 2
        try:
            accounts = read_accounts(chosen)
        except (OSError, ValueError, openpyxl.utils.exceptions.InvalidFileException) as exc:
            print(f"Lỗi đọc tài khoản: {exc}", file=sys.stderr)
            return 2
        source = str(chosen.resolve())
    if args.limit > 0:
        accounts = accounts[: args.limit]
    if not accounts:
        print("Không có tài khoản để cào.", file=sys.stderr)
        return 2

    workers = max(1, min(args.workers, len(accounts)))
    print(f"Nguồn: {source}")
    print(f"Số tài khoản: {len(accounts)} | Luồng: {workers} | Thử lại: {args.retries}")
    print(f"Excel: {output}")
    print(f"Nhật ký: {log_path}")
    try:
        results = scan_all(accounts, workers, max(0, args.retries), log_path)
        saved = build_excel(results, output)
        write_schema(log_path.with_name("schema.json"), _model_links or [])
    except KeyboardInterrupt:
        print("\nĐã dừng.", file=sys.stderr)
        return 130
    success = sum(item.scan_status == "Thành công" for item in results)
    partial = sum(item.scan_status == "Lỗi một phần" for item in results)
    missing = sum(item.scan_status == "Chưa có hồ sơ" for item in results)
    denied = sum(item.scan_status == "Sai tài khoản/mật khẩu" for item in results)
    failed = len(results) - success - partial - missing - denied
    print(f"Xong: thành công={success}, một phần={partial}, chưa có hồ sơ={missing}, sai mật khẩu={denied}, lỗi={failed}")
    print(f"File: {saved}")
    print(f"Nhật ký: {log_path}")
    return 0 if failed == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
