"""Highlight dòng MSSV trên PDF danh sách (scan) theo data.xlsx và đơn vị KHKTMT."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import sys
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import cv2
import numpy as np
import pymupdf
from openpyxl import load_workbook
from rapidocr_onnxruntime import RapidOCR

HERE = Path(__file__).resolve().parent
DEFAULT_ROOT = HERE
DEFAULT_DATA = HERE / "data.xlsx"
DEFAULT_OUTPUT = HERE / "highlighted"
DEFAULT_LOG = HERE / "highlight.log"
DEFAULT_FACULTY = "Khoa học và Kỹ thuật Máy tính"
DEFAULT_RENDER_SCALE = 2.0

MSSV_RE = re.compile(r"^\d{7}$")
INVALID_WIN_CHARS = re.compile(r'[<>:"/\\|?*]')

# Thư mục con chứa PDF nguồn (bỏ qua output / cache / probe)
SKIP_DIR_NAMES = frozenset({"highlighted", ".ocr_cache", "_probe"})


@dataclass
class StudentRef:
    mssv: str
    full_name: str  # họ tên có dấu từ data.xlsx


@dataclass
class RowHit:
    mssv: str
    page_index: int
    rect_pdf: tuple[float, float, float, float]


@dataclass
class RunConfig:
    log_file: Path | None
    quiet: bool


def configure_stdout() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def log_line(cfg: RunConfig, status: str, src: str, detail: str, *, echo: bool = True) -> None:
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    line = f"{ts}\t{status}\t{src}\t{detail}\n"
    if cfg.log_file is not None:
        with cfg.log_file.open("a", encoding="utf-8") as f:
            f.write(line)
    if echo and not cfg.quiet:
        print(line.rstrip())


def strip_accents(text: str) -> str:
    text = unicodedata.normalize("NFD", text)
    text = "".join(c for c in text if unicodedata.category(c) != "Mn")
    return text.replace("đ", "d").replace("Đ", "D")


def norm_key(text: str) -> str:
    return strip_accents(text).lower().strip()


def is_target_faculty(text: str, faculty_norm: str) -> bool:
    n = norm_key(text)
    if faculty_norm in n:
        return True
    return "ky thuat may tinh" in n


def sanitize_filename_part(text: str) -> str:
    text = INVALID_WIN_CHARS.sub("", text.strip())
    return text.rstrip(". ") or "unknown"


def output_basename(mssv: str, folder_name: str, full_name: str, index: int) -> str:
    if index <= 1:
        middle = folder_name
    else:
        middle = f"{folder_name}-{index}"
    return f"{mssv}_{middle}_{full_name}.pdf"


def box_center(box: list[list[float]]) -> tuple[float, float]:
    xs = [p[0] for p in box]
    ys = [p[1] for p in box]
    return sum(xs) / 4, sum(ys) / 4


def box_rect(box: list[list[float]]) -> tuple[float, float, float, float]:
    xs = [p[0] for p in box]
    ys = [p[1] for p in box]
    return min(xs), min(ys), max(xs), max(ys)


def group_rows(
    items: list[tuple[list[list[float]], str, float]],
    *,
    y_tol: float,
) -> list[list[tuple[list[list[float]], str, float]]]:
    items = sorted(items, key=lambda x: box_center(x[0])[1])
    rows: list[dict[str, object]] = []
    for it in items:
        cy = box_center(it[0])[1]
        for row in rows:
            if abs(cy - float(row["y"])) <= y_tol:  # type: ignore[arg-type]
                row_items = row["items"]
                assert isinstance(row_items, list)
                row_items.append(it)
                row["y"] = sum(box_center(x[0])[1] for x in row_items) / len(row_items)
                break
        else:
            rows.append({"y": cy, "items": [it]})
    return [row["items"] for row in rows]  # type: ignore[return-value]


def parse_table_row(
    row_items: list[tuple[list[list[float]], str, float]],
    *,
    faculty_norm: str,
) -> tuple[str, tuple[float, float, float, float]] | None:
    cells: list[tuple[float, tuple[float, float, float, float], str]] = []
    for box, text, _score in row_items:
        rect = box_rect(box)
        cells.append((rect[0], rect, text.strip()))
    cells.sort(key=lambda c: c[0])

    mssv: str | None = None
    has_unit = False

    for _x0, _rect, t in cells:
        compact = t.replace(" ", "")
        if MSSV_RE.match(compact):
            mssv = compact
        elif is_target_faculty(t, faculty_norm):
            has_unit = True

    if not mssv or not has_unit:
        return None

    rects = [box_rect(box) for box, _text, _score in row_items]
    xs = [r[0] for r in rects]
    ys = [r[1] for r in rects]
    x2s = [r[2] for r in rects]
    y2s = [r[3] for r in rects]
    row_rect = (min(xs), min(ys), max(x2s), max(y2s))
    return mssv, row_rect


def pixmap_to_bgr(pix: pymupdf.Pixmap) -> np.ndarray:
    img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.h, pix.w, pix.n)
    if pix.n == 4:
        return cv2.cvtColor(img, cv2.COLOR_RGBA2BGR)
    if pix.n == 3:
        return cv2.cvtColor(img, cv2.COLOR_RGB2BGR)
    return img


def pdf_rect_from_pix_rect(
    rect_px: tuple[float, float, float, float],
    *,
    scale: float,
) -> pymupdf.Rect:
    inv = 1.0 / scale
    return pymupdf.Rect(rect_px[0] * inv, rect_px[1] * inv, rect_px[2] * inv, rect_px[3] * inv)


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def cache_path_for(pdf_path: Path, cache_dir: Path) -> Path:
    digest = file_sha256(pdf_path)
    return cache_dir / f"{pdf_path.stem}_{digest[:16]}.json"


def load_ocr_cache(path: Path) -> list[RowHit] | None:
    if not path.is_file():
        return None
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return None
    hits: list[RowHit] = []
    for item in raw.get("rows", []):
        hits.append(
            RowHit(
                mssv=item["mssv"],
                page_index=int(item["page_index"]),
                rect_pdf=tuple(item["rect_pdf"]),
            )
        )
    return hits


def save_ocr_cache(path: Path, hits: list[RowHit], *, meta: dict[str, object]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "meta": meta,
        "rows": [
            {
                "mssv": h.mssv,
                "page_index": h.page_index,
                "rect_pdf": list(h.rect_pdf),
            }
            for h in hits
        ],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def ocr_pdf_rows(
    pdf_path: Path,
    *,
    faculty: str,
    scale: float,
    y_tol: float,
    ocr: RapidOCR,
) -> list[RowHit]:
    faculty_norm = norm_key(faculty)
    doc = pymupdf.open(pdf_path)
    matrix = pymupdf.Matrix(scale, scale)
    hits: list[RowHit] = []
    try:
        for page_index in range(doc.page_count):
            page = doc[page_index]
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            img = pixmap_to_bgr(pix)
            result, _elapsed = ocr(img)
            if not result:
                continue
            for row_items in group_rows(result, y_tol=y_tol):
                parsed = parse_table_row(row_items, faculty_norm=faculty_norm)
                if parsed is None:
                    continue
                mssv, rect_px = parsed
                rect_pdf = pdf_rect_from_pix_rect(rect_px, scale=scale)
                hits.append(
                    RowHit(
                        mssv=mssv,
                        page_index=page_index,
                        rect_pdf=(rect_pdf.x0, rect_pdf.y0, rect_pdf.x1, rect_pdf.y1),
                    )
                )
    finally:
        doc.close()
    return hits


def load_students_xlsx(path: Path) -> dict[str, StudentRef]:
    if not path.is_file():
        raise FileNotFoundError(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            return {}
        headers = [str(c).strip() if c is not None else "" for c in header]
        lower = [h.lower() for h in headers]

        def col_index(*candidates: str) -> int | None:
            for cand in candidates:
                cand_l = cand.lower()
                for i, h in enumerate(lower):
                    if cand_l in h or h in cand_l:
                        return i
            return None

        i_mssv = col_index("mssv", "ma so sinh vien", "mã số sinh viên")
        if i_mssv is None:
            raise ValueError("Không tìm thấy cột MSSV trong data.xlsx")
        i_name = col_index("ho ten", "họ tên", "ho va ten", "ten", "tên")
        if i_name is None:
            raise ValueError("Không tìm thấy cột Họ tên trong data.xlsx")

        out: dict[str, StudentRef] = {}
        for row in rows:
            if not row or i_mssv >= len(row):
                continue
            raw_mssv = row[i_mssv]
            if raw_mssv is None:
                continue
            if isinstance(raw_mssv, float) and raw_mssv.is_integer():
                mssv = str(int(raw_mssv))
            else:
                mssv = str(raw_mssv).strip().replace(" ", "")
            if not MSSV_RE.match(mssv):
                continue
            if i_name >= len(row) or row[i_name] is None:
                continue
            full_name = str(row[i_name]).strip()
            if not full_name:
                continue
            out[mssv] = StudentRef(mssv=mssv, full_name=full_name)
        return out
    finally:
        wb.close()


def collect_pdfs(root: Path, *, only_folders: set[str] | None) -> list[tuple[Path, str]]:
    """Trả về (đường dẫn PDF, tên folder loại vd. TNTT)."""
    items: list[tuple[Path, str]] = []
    for category_dir in sorted(root.iterdir()):
        if not category_dir.is_dir():
            continue
        if category_dir.name in SKIP_DIR_NAMES:
            continue
        folder_name = category_dir.name
        if only_folders is not None and folder_name not in only_folders:
            continue
        pdfs = sorted(category_dir.glob("*.pdf"))
        for pdf in pdfs:
            items.append((pdf, folder_name))
    return items


def draw_highlight(page: pymupdf.Page, rect: pymupdf.Rect) -> None:
    shape = page.new_shape()
    shape.draw_rect(rect)
    shape.finish(color=(1, 1, 0), fill=(1, 1, 0), fill_opacity=0.35, width=0)
    shape.commit()


def write_highlighted_pdf(
    src_pdf: Path,
    dest_pdf: Path,
    row: RowHit,
) -> None:
    doc = pymupdf.open(src_pdf)
    try:
        page = doc[row.page_index]
        rect = pymupdf.Rect(*row.rect_pdf)
        draw_highlight(page, rect)
        dest_pdf.parent.mkdir(parents=True, exist_ok=True)
        doc.save(dest_pdf)
    finally:
        doc.close()


def run(
    root: Path,
    data_path: Path,
    output_root: Path,
    cfg: RunConfig,
    *,
    faculty: str,
    scale: float,
    y_tol: float,
    dry_run: bool,
    skip_existing: bool,
    limit: int,
    only_folders: set[str] | None,
    use_cache: bool,
    refresh_cache: bool,
) -> int:
    students = load_students_xlsx(data_path)
    if not students:
        log_line(cfg, "ERROR", str(data_path), "Danh sách MSSV trong data.xlsx trống")
        return 2

    pdfs = collect_pdfs(root, only_folders=only_folders)
    if not pdfs:
        log_line(cfg, "ERROR", str(root), "Không tìm thấy PDF trong folder con (vd. TNTT/TNTT.pdf)")
        return 2

    cache_dir = root / ".ocr_cache"
    ocr = RapidOCR()
    ok = 0
    skip = 0
    miss = 0
    errors = 0

    for pdf_path, folder_name in pdfs:
        cache_file = cache_path_for(pdf_path, cache_dir)
        hits: list[RowHit] | None = None
        if use_cache and not refresh_cache:
            hits = load_ocr_cache(cache_file)
            if hits is not None:
                log_line(cfg, "CACHE", str(pdf_path), f"{len(hits)} dòng từ .ocr_cache")

        if hits is None:
            log_line(cfg, "OCR", str(pdf_path), f"scale={scale} faculty={faculty}")
            if dry_run:
                log_line(cfg, "SKIP", str(pdf_path), "dry-run: bỏ qua OCR thật")
                continue
            hits = ocr_pdf_rows(
                pdf_path,
                faculty=faculty,
                scale=scale,
                y_tol=y_tol,
                ocr=ocr,
            )
            if use_cache:
                save_ocr_cache(
                    cache_file,
                    hits,
                    meta={
                        "pdf": str(pdf_path),
                        "faculty": faculty,
                        "scale": scale,
                        "y_tol": y_tol,
                    },
                )
            log_line(cfg, "OCR_OK", str(pdf_path), f"{len(hits)} dòng đơn vị khớp")

        by_mssv: dict[str, RowHit] = {}
        for h in hits:
            by_mssv[h.mssv] = h

        targets = [s for mssv, s in students.items() if mssv in by_mssv]
        if limit > 0:
            targets = targets[:limit]

        for student in targets:
            row = by_mssv.get(student.mssv)
            if row is None:
                miss += 1
                log_line(
                    cfg,
                    "MISS_PDF",
                    student.mssv,
                    f"Không thấy dòng KHKTMT trong {pdf_path.name}",
                )
                continue

            display_name = sanitize_filename_part(student.full_name)

            dest: Path | None = None
            for index in range(1, 1000):
                name = output_basename(student.mssv, folder_name, display_name, index)
                candidate = output_root / name
                if candidate.exists() and skip_existing:
                    skip += 1
                    log_line(cfg, "SKIP", name, "đã tồn tại")
                    dest = None
                    break
                if not candidate.exists() or not skip_existing:
                    dest = candidate
                    break

            if dest is None:
                continue

            if dry_run:
                ok += 1
                log_line(cfg, "DRY", dest.name, f"page={row.page_index + 1}")
                continue

            try:
                write_highlighted_pdf(pdf_path, dest, row)
                ok += 1
                log_line(cfg, "OK", dest.name, f"page={row.page_index + 1}")
            except OSError as exc:
                errors += 1
                log_line(cfg, "ERROR", dest.name, str(exc))

    log_line(
        cfg,
        "SUMMARY",
        str(output_root),
        f"ok={ok} skip={skip} miss={miss} err={errors}",
        echo=True,
    )
    return 1 if errors else 0


def main() -> None:
    configure_stdout()
    parser = argparse.ArgumentParser(description="Highlight dòng MSSV trên PDF danh sách scan.")
    parser.add_argument(
        "--root",
        type=Path,
        default=DEFAULT_ROOT,
        help="Thư mục highlighter (chứa data.xlsx và folder loại PDF)",
    )
    parser.add_argument(
        "--data",
        type=Path,
        default=DEFAULT_DATA,
        help="File Excel danh sách MSSV (mặc định: highlighter/data.xlsx)",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help="Thư mục PDF đã highlight (mặc định: highlighter/highlighted/)",
    )
    parser.add_argument(
        "--faculty",
        default=DEFAULT_FACULTY,
        help="Chỉ highlight dòng thuộc đơn vị này (so khớp không dấu)",
    )
    parser.add_argument(
        "--scale",
        type=float,
        default=DEFAULT_RENDER_SCALE,
        help="Tỉ lệ render OCR (2.0 thường ổn với PDF scan A4)",
    )
    parser.add_argument(
        "--y-tol",
        type=float,
        default=18.0,
        help="Ngưỡng ghép ô OCR cùng hàng (pixel theo scale render)",
    )
    parser.add_argument(
        "--only",
        action="append",
        metavar="FOLDER",
        help="Chỉ xử lý folder loại (vd. --only TNTT)",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=0,
        help="Chỉ xử lý N MSSV đầu sau lọc (0 = tất cả)",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Không ghi PDF; vẫn log kết quả dự kiến",
    )
    parser.add_argument(
        "--skip-existing",
        action="store_true",
        help="Bỏ qua nếu file đích đã tồn tại",
    )
    parser.add_argument(
        "--no-cache",
        action="store_true",
        help="Không đọc/ghi cache OCR (.ocr_cache)",
    )
    parser.add_argument(
        "--refresh-cache",
        action="store_true",
        help="OCR lại và cập nhật cache",
    )
    parser.add_argument(
        "--clean-output",
        action="store_true",
        help="Xóa thư mục output trước khi chạy (không áp dụng khi --dry-run)",
    )
    parser.add_argument(
        "--quiet",
        action="store_true",
        help="Không in từng dòng ra console",
    )
    parser.add_argument(
        "--no-log",
        action="store_true",
        help="Không ghi highlight.log",
    )
    parser.add_argument(
        "--log",
        type=Path,
        default=None,
        help="Đường dẫn file log tùy chọn",
    )
    args = parser.parse_args()

    root = args.root.resolve()
    data_path = args.data.resolve()
    output_root = args.output.resolve()

    if args.clean_output and not args.dry_run:
        if output_root.exists():
            shutil.rmtree(output_root)

    only_folders: set[str] | None = None
    if args.only:
        only_folders = set(args.only)

    log_path: Path | None = None
    if not args.no_log:
        log_path = (args.log or DEFAULT_LOG).resolve()

    cfg = RunConfig(log_file=log_path, quiet=args.quiet)

    try:
        code = run(
            root,
            data_path,
            output_root,
            cfg,
            faculty=args.faculty,
            scale=max(0.5, args.scale),
            y_tol=max(4.0, args.y_tol),
            dry_run=args.dry_run,
            skip_existing=args.skip_existing,
            limit=max(0, args.limit),
            only_folders=only_folders,
            use_cache=not args.no_cache,
            refresh_cache=args.refresh_cache,
        )
    except FileNotFoundError as exc:
        print(f"Không tìm thấy file: {exc}", file=sys.stderr)
        raise SystemExit(2) from exc
    except ValueError as exc:
        print(str(exc), file=sys.stderr)
        raise SystemExit(2) from exc

    raise SystemExit(code)


if __name__ == "__main__":
    main()
